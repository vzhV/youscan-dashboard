from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time

from openpyxl.utils import get_column_letter

from app.core.errors import InvalidFileError
from app.importing.limits import ImportLimits, too_many_columns, too_many_rows
from app.importing.location import describe_source

Cell = str | float | datetime | None


@dataclass(frozen=True)
class Row:
    number: int  # 1-based spreadsheet row, so errors can point the user at the exact line
    cells: tuple[Cell, ...]


@dataclass(frozen=True)
class Table:
    name: str | None  # None for CSV, which has no sheets
    header: tuple[str, ...]
    rows: tuple[Row, ...]


def normalize_cell(value: object) -> Cell:
    if value is None:
        return None
    # bool must precede the number check because bool is a subclass of int.
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return str(value).strip() or None


def build_table(
    name: str | None,
    raw_rows: Iterable[tuple[int, Sequence[object]]],
    limits: ImportLimits,
) -> Table | None:
    rows = _non_blank_rows(name, raw_rows, limits)
    header_row = next(rows, None)
    if header_row is None:
        return None
    data_rows = tuple(rows)
    kept = _kept_columns(name, header_row, data_rows)
    header = tuple(_header_text(_cell_at(header_row, index)) for index in kept)
    _reject_duplicate_headers(name, header, kept)
    return Table(name, header, tuple(_select(row, kept) for row in data_rows))


def _non_blank_rows(
    name: str | None, raw_rows: Iterable[tuple[int, Sequence[object]]], limits: ImportLimits
) -> Iterator[Row]:
    # Limits are checked while streaming so an oversized sheet is rejected before it is held in memory.
    data_row_count = -1  # the first non-blank row is the header, not data
    for number, raw in raw_rows:
        cells = _trim(tuple(normalize_cell(value) for value in raw))
        if not cells:
            continue
        data_row_count += 1
        if len(cells) > limits.max_columns:
            raise too_many_columns(name, limits.max_columns, len(cells), number)
        if data_row_count > limits.max_rows:
            raise too_many_rows(name, limits.max_rows, data_row_count, number)
        yield Row(number, cells)


def _trim(cells: tuple[Cell, ...]) -> tuple[Cell, ...]:
    # Read-only sheets pad rows to the declared width; trailing empties carry no data.
    end = len(cells)
    while end and cells[end - 1] is None:
        end -= 1
    return cells[:end]


def _cell_at(row: Row, index: int) -> Cell:
    # Rows are ragged after trimming, so a missing trailing cell counts as empty.
    return row.cells[index] if index < len(row.cells) else None


def _select(row: Row, columns: Sequence[int]) -> Row:
    return Row(row.number, tuple(_cell_at(row, index) for index in columns))


def _kept_columns(name: str | None, header_row: Row, data_rows: Sequence[Row]) -> list[int]:
    header_columns = {index for index, cell in enumerate(header_row.cells) if cell is not None}
    data_columns = {index for row in data_rows for index, cell in enumerate(row.cells) if cell is not None}
    orphans = data_columns - header_columns
    if orphans:
        letter = get_column_letter(min(orphans) + 1)
        raise InvalidFileError(
            f"{describe_source(name)} has data in column {letter} but no header",
            details={"sheet": name, "column": letter},
        )
    return sorted(header_columns)


def _reject_duplicate_headers(name: str | None, header: Sequence[str], columns: Sequence[int]) -> None:
    seen: set[str] = set()
    for title, index in zip(header, columns, strict=True):
        if title in seen:
            raise InvalidFileError(
                f"{describe_source(name)} has a duplicate column '{title}'",
                details={"sheet": name, "column": get_column_letter(index + 1)},
            )
        seen.add(title)


def _header_text(cell: Cell) -> str:
    # XLSX stores every number as a float; "2023.0" or a midnight timestamp would be an ugly header.
    if isinstance(cell, datetime):
        return cell.date().isoformat() if cell.time() == time.min else cell.isoformat()
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell)

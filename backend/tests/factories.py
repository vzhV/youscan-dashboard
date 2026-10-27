from collections.abc import Mapping, Sequence
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook

from app.core.enums import ChartType, DatasetOrigin
from app.importing.reader import read_tables
from app.importing.source_file import SourceFile
from app.importing.table import Cell, Row, Table
from app.models import Dataset

SEED_DATA_DIR = Path(__file__).resolve().parents[1] / "seed_data"
FIRST_DATA_ROW = 2  # row 1 is the header


def make_xlsx(sheets: Mapping[str, Sequence[Sequence[object]]]) -> bytes:
    workbook = Workbook(write_only=True)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_xlsx_with_cells(
    name: str, rows: Sequence[Sequence[object]], extra_cells: Mapping[str, str | float]
) -> bytes:
    """Single-sheet workbook with cells set by coordinate.

    Unlike make_xlsx this is not write-only, so the file declares a <dimension> covering the
    far-away cells, which is what makes read-only rows come back padded.
    """
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = name
    for row in rows:
        sheet.append(list(row))
    for coordinate, value in extra_cells.items():
        sheet[coordinate] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_zip(members: Mapping[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for member, data in members.items():
            archive.writestr(member, data)
    return buffer.getvalue()


def replace_zip_member(archive_bytes: bytes, member: str, data: bytes) -> bytes:
    with ZipFile(BytesIO(archive_bytes)) as archive:
        members = {item: archive.read(item) for item in archive.namelist()}
    members[member] = data
    return make_zip(members)


def seed_table(filename: str, sheet: str | None = None) -> Table:
    """Parse a provided sample file; `sheet` picks one table from a workbook."""
    source = SourceFile(filename, (SEED_DATA_DIR / filename).read_bytes())
    tables = [table for table in read_tables(source) if sheet is None or table.name == sheet]
    if not tables:
        raise LookupError(f"{filename} has no sheet named '{sheet}'")
    return tables[0]


def make_table(header: Sequence[str], rows: Sequence[Sequence[Cell]], name: str | None = None) -> Table:
    """Like build_table, short rows are padded to the header width."""
    width = len(header)
    return Table(
        name,
        tuple(header),
        tuple(
            Row(number, (*cells, *([None] * (width - len(cells)))))
            for number, cells in enumerate(rows, start=FIRST_DATA_ROW)
        ),
    )


def make_dataset() -> Dataset:
    return Dataset(
        name="data",
        origin=DatasetOrigin.SEED,
        chart_type=ChartType.LINE,
        row_count=1,
        data={"categories": [], "series": []},
        display_data={"categories": [], "series": []},
    )
